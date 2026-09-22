"""Write the league's books to a workbook."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import prizes as pz
from . import treasury as tr

KRW = '#,##0 "KRW"'
INK = "1F2A17"
HEAD_FILL = PatternFill("solid", fgColor="14301F")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE = Font(bold=True, size=13, color=INK)
MUTED = Font(size=9, color="6B7280")
EDGE = Side(style="thin", color="D8D8CE")
BORDER = Border(bottom=EDGE)


def _header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 20
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _money(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    c.number_format = KRW
    return c


def build(snapshot, paid_ids, path):
    ledger = snapshot["prizes"]
    books = tr.summary(ledger, paid_ids)
    wb = Workbook()

    # --- Summary ---------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{snapshot['league']['name']} {2026}/27 - prize pool"
    ws["A1"].font = TITLE
    ws["A2"] = f"Snapshot {snapshot['generated'][:16].replace('T', ' ')} UTC"
    ws["A2"].font = MUTED
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18

    rows = [
        ("MONEY IN", None),
        (f"Entries ({books['players']} x {books['entry']:,} KRW)", books["collected"]),
        ("", None),
        ("MONEY OUT", None),
        ("Paid to weekly winners", books["paidOut"]),
        ("Due now (weeks won, not yet paid)", books["dueNow"]),
        ("", None),
        ("STILL RESERVED", None),
        ("Future weekly prizes", books["reservedWeekly"]),
        ("Special awards", books["reservedSpecials"]),
        ("Champion and runner-up", books["reservedPodium"]),
        ("", None),
        ("POSITION", None),
        ("Balance on hand", books["balance"]),
        ("Total still owed", books["committed"]),
        ("Unallocated (should be zero)", books["unallocated"]),
    ]
    r = 4
    for label, value in rows:
        if value is None:
            if label:
                c = ws.cell(row=r, column=1, value=label)
                c.font = Font(bold=True, size=9, color="6B7280")
            r += 1
            continue
        ws.cell(row=r, column=1, value=label)
        _money(ws, r, 2, value)
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=2).border = BORDER
        r += 1

    ws.cell(row=r + 1, column=1,
            value=f"{books['paidCount']} of {books['totalCount']} weekly prizes paid")
    ws.cell(row=r + 1, column=1).font = MUTED

    # --- Prize ledger ----------------------------------------------------
    ws = wb.create_sheet("Prize ledger")
    ws["A1"] = "Weekly prizes"
    ws["A1"].font = TITLE
    _header(ws, 3, ["GW", "Winner", "Score", "Amount", "Status", "Tie"],
            [6, 24, 8, 16, 12, 8])
    r = 4
    for line in books["lines"]:
        ws.cell(row=r, column=1, value=line["gw"])
        ws.cell(row=r, column=2, value=line["team"])
        ws.cell(row=r, column=3, value=line["points"])
        _money(ws, r, 4, line["amount"])
        ws.cell(row=r, column=5, value="PAID" if line["paid"] else "DUE")
        ws.cell(row=r, column=6, value="split" if line["tied"] else "")
        if not line["paid"]:
            ws.cell(row=r, column=5).font = Font(bold=True, color="A83246")
        r += 1
    ws.auto_filter.ref = f"A3:F{r - 1}"
    ws.cell(row=r + 1, column=3, value="Total")
    ws.cell(row=r + 1, column=3).font = Font(bold=True)
    _money(ws, r + 1, 4, sum(l["amount"] for l in books["lines"])).font = Font(bold=True)

    # --- Per manager -----------------------------------------------------
    ws = wb.create_sheet("By manager")
    ws["A1"] = "What each manager has received"
    ws["A1"].font = TITLE
    _header(ws, 3, ["Pos", "Team", "Manager", "Entry paid", "Received", "Still owed"],
            [6, 24, 22, 16, 16, 16])
    people = {m["entry"]: m["manager"] for m in snapshot["managers"]}
    r = 4
    for row in tr.by_manager(books, snapshot["standings"]):
        ws.cell(row=r, column=1, value=row["rank"])
        ws.cell(row=r, column=2, value=row["team"])
        ws.cell(row=r, column=3, value=people.get(row["entry"], ""))
        _money(ws, r, 4, row["entryFee"])
        _money(ws, r, 5, row["received"])
        _money(ws, r, 6, row["owed"])
        r += 1

    # --- Gameweeks -------------------------------------------------------
    ws = wb.create_sheet("Gameweeks")
    ws["A1"] = "Every score, every gameweek"
    ws["A1"].font = TITLE
    _header(ws, 3,
            ["GW", "Pos", "Team", "Manager", "GW pts", "Gross", "Hit",
             "Bench", "Captain", "C pts", "Season total", "Chip"],
            [6, 6, 24, 22, 9, 9, 7, 8, 16, 8, 14, 12])
    players = snapshot["players"]
    r = 4
    for gw in snapshot["playedGws"]:
        season = {x["entry"]: x for x in snapshot["season"][str(gw)]}
        for row in snapshot["gameweeks"][str(gw)]["rows"]:
            cap = players.get(str(row["captain"]), {})
            vals = [
                gw, row["place"], row["team"], row["manager"], row["points"],
                row["gross"], -row["hit"] if row["hit"] else 0, row["bench"],
                cap.get("name", ""), row["captainPoints"],
                season.get(row["entry"], {}).get("total"), row["chip"] or "",
            ]
            for i, v in enumerate(vals, start=1):
                ws.cell(row=r, column=i, value=v)
            r += 1
    ws.auto_filter.ref = f"A3:L{r - 1}"

    # --- Special awards --------------------------------------------------
    ws = wb.create_sheet("Special awards")
    ws["A1"] = "Special awards - provisional"
    ws["A1"].font = TITLE
    _header(ws, 3, ["Award", "Amount", "Currently", "Detail"], [30, 16, 24, 30])
    e = ledger["extremes"]
    fifth = next((s for s in snapshot["standings"] if s["rank"] == 5), None)
    awards = [
        ("Highest single gameweek", pz.SPECIALS["highest"],
         e["highest"]["team"] if e["highest"] else "-",
         f"{e['highest']['points']} in GW{e['highest']['gw']}" if e["highest"] else ""),
        ("Exactly 111 in a gameweek", pz.SPECIALS["exact111"],
         ", ".join(x["team"] for x in e["exact111"]) or "unclaimed", ""),
        ("Lowest single gameweek", pz.SPECIALS["lowest"],
         e["lowest"]["team"] if e["lowest"] else "-",
         f"{e['lowest']['points']} in GW{e['lowest']['gw']}, before hits" if e["lowest"] else ""),
        ("5th at season end", pz.SPECIALS["fifth"],
         fifth["team"] if fifth else "-", "currently 5th"),
        ("Champion", pz.CHAMPION, snapshot["standings"][0]["team"], "currently 1st"),
        ("Runner-up", pz.RUNNER_UP, snapshot["standings"][1]["team"], "currently 2nd"),
    ]
    r = 4
    for name, amount, who, detail in awards:
        ws.cell(row=r, column=1, value=name)
        _money(ws, r, 2, amount)
        ws.cell(row=r, column=3, value=who)
        ws.cell(row=r, column=4, value=detail)
        r += 1
    ws.cell(row=r + 1, column=1,
            value="Provisional: decided at the end of the season.").font = MUTED

    wb.save(path)
    return books
